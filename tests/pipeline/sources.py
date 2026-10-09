"""Source documents built inside the tests (PDF via PyMuPDF, DOCX via python-docx) + sample text."""

from __future__ import annotations

import io

from tests.pipeline.fakes import png_bytes

SAMPLE_MARKDOWN = """<!-- page 1 -->

# Ohm's Law

## Voltage and current

Voltage is the electrical pressure that pushes charge around a circuit. It is measured in volts.
Current is the rate at which charge flows past a point in the circuit. It is measured in amperes.
A higher voltage across the same wire pushes a larger current through it.

[Figure fig-p1-1: Figure 1: A simple circuit with a cell and a resistor]

<!-- page 2 -->

## Resistance

Resistance is the opposition a material offers to the flow of current. It is measured in ohms.
Copper wires have a low resistance while rubber has a very high resistance.
Long, thin and hot conductors resist the current more than short, thick and cold ones.

## Ohm's law

Ohm's law states that the current through a conductor is proportional to the voltage across it.
In symbols, V = I × R where V is the voltage, I is the current and R is the resistance.
The law holds for metallic conductors at a constant temperature.

| Quantity | Symbol | Unit |
|---|---|---|
| Voltage | V | volt |
| Current | I | ampere |

<!-- page 3 -->

## Power in circuits

Electrical power is the rate at which energy is converted in a circuit. It is measured in watts.
The power used by a device is P = V × I for any device in the circuit.
A geyser converts electrical energy into heat at a much higher rate than an LED lamp.
"""


def make_pdf() -> bytes:
    """3-page PDF: running header, page numbers, headings, an embedded PNG figure + caption, a ruled table."""
    import pymupdf

    doc = pymupdf.open()
    figure = png_bytes(320, 200, "circuit")
    body = {
        1: ("1. Voltage and current", [
            "Voltage is the electrical pressure that pushes charge around a circuit.",
            "Current is the rate at which charge flows past a point in the circuit.",
            "A higher voltage across the same wire pushes a larger current through it.",
        ]),
        2: ("2. Resistance", [
            "Resistance is the opposition a material offers to the flow of current.",
            "Copper wires have a low resistance while rubber has a very high resistance.",
            "Ohm's law states that V = I x R for metallic conductors.",
        ]),
        3: ("3. Units", [
            "Each electrical quantity has its own unit and symbol.",
            "The table below lists the most common ones.",
        ]),
    }
    for n in (1, 2, 3):
        page = doc.new_page()
        page.insert_text((72, 40), "Basic Electrical Engineering - Unit 2", fontsize=8)
        title, lines = body[n]
        page.insert_text((72, 90), title, fontsize=20)
        y = 130.0
        for line in lines:
            page.insert_text((72, y), line, fontsize=11)
            y += 18
        if n == 2:
            page.insert_image(pymupdf.Rect(72, 260, 392, 460), stream=figure)
            page.insert_text((72, 480), "Figure 1: A resistor connected to a cell", fontsize=10)
        if n == 3:
            rows = [["Quantity", "Symbol", "Unit"], ["Voltage", "V", "volt"], ["Current", "I", "ampere"],
                    ["Resistance", "R", "ohm"]]
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    rect = pymupdf.Rect(72 + c * 120, 260 + r * 24, 72 + (c + 1) * 120, 260 + (r + 1) * 24)
                    page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                    page.insert_text((rect.x0 + 4, rect.y0 + 16), cell, fontsize=10)
        page.insert_text((300, 820), f"Page {n}", fontsize=8)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    """DOCX with a heading, paragraphs, a bullet list, a table and an embedded image + caption."""
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("Ohm's Law", level=1)
    d.add_paragraph("Ohm's law relates the voltage across a conductor to the current through it.")
    d.add_heading("Resistance", level=2)
    p = d.add_paragraph("Resistance is the ")
    p.add_run("opposition").bold = True
    p.add_run(" a material offers to current.")
    d.add_paragraph("Copper has low resistance", style="List Bullet")
    d.add_paragraph("Rubber has high resistance", style="List Bullet")
    d.add_picture(io.BytesIO(png_bytes(300, 200, "docx-figure")), width=Inches(3))
    d.add_paragraph("Figure 2: Resistors in a circuit")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Quantity"
    table.cell(0, 1).text = "Unit"
    table.cell(1, 0).text = "Resistance"
    table.cell(1, 1).text = "ohm"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
