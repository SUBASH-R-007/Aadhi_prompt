"""An SME video script end to end with the offline responders: what the model sees, and the lecture it makes.

* The SME's name, the reviewer, the video duration and the course code never reach a prompt, the
  screenplay or a job event, and lint finds no ``content.admin_leak``.
* The two-clip script becomes one continuous lecture: one opening that states the objectives once, no
  per-clip title cards or "bridge to next part" scenes, concepts in source order, a ``bridge_in`` for
  every scene after the first, and quizzes built from the source's own questions.
* ``fake_content`` reads sources the way the offline demo needs (sentences joined across script line
  breaks, section roles, questions with answers, formulas).
"""

from __future__ import annotations

import asyncio
import io
import json
import re
from typing import Any

import pytest

from aadhi.pipeline import fake_content, orchestrator
from aadhi.pipeline.base import GenerationOptions
from aadhi.pipeline.brief import brief_problems, build_brief_prompt
from aadhi.pipeline.chunking import chunk_markdown
from aadhi.pipeline.docx_extract import docx_to_markdown
from aadhi.pipeline.gen_models import GenBrief, GenPlan
from aadhi.pipeline.prompting import extract_json
from aadhi.pipeline.source_scope import scope_source
from aadhi.pipeline.validate import lint_intros
from aadhi.schemas.screenplay import QuizScene, Screenplay
from tests.pipeline.dbutil import get_version, seed_project
from tests.pipeline.fakes import install_providers
from tests.pipeline.fixtures.sme_sources import sample_template_bytes

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OPTIONS = GenerationOptions(target_minutes=8, quiz_every_n_concepts=2)

SME = "Dr. Ramesh Kumar"
REVIEWER = "Prof. Anitha"
DURATION = "12 minutes"
COURSE_CODE = "EE3251"
# every form of the removed header values a leak could take
SECRETS = ("Ramesh Kumar", "Ramesh", "Anitha", DURATION, COURSE_CODE, "14/08/2025")
PACKAGING = ("Estimated duration", "CLIP 1 SCRIPT", "CLIP 2 SCRIPT", "TITLE CARD", "BRIDGE TO NEXT PART", "COMING UP NEXT",
             "Aadhi speaks:", "BOARD displays", "[0:", "fade to black", "Thank you for watching")

HEADER = [
    f"SME Name: {SME}",
    "Designation: Assistant Professor",
    "Department: Electrical and Electronics Engineering",
    f"Course Code: {COURSE_CODE}",
    f"Video Duration: {DURATION}",
    f"Reviewed by: {REVIEWER}",
    "Date: 14/08/2025",
]
# (style, text): "h" = Heading 2 (as SME templates mark clip / segment lines), "p" = paragraph
SCRIPT: list[tuple[str, str]] = [
    ("h", "SUBJECT NAME : BASIC ELECTRICAL ENGINEERING"),
    ("h", "UNIT NAME : DC CIRCUITS"),
    ("p", "SESSION 3 - Series and Parallel Resistors"),
    ("h", "CLIP 1 SCRIPT - SERIES RESISTORS"),
    ("p", "Estimated duration: 4 minutes"),
    ("h", "SEGMENT 1 - WHAT THIS VIDEO WILL COVER [0:00 - 0:30]"),
    ("p", "BOARD Title: LEARNING OBJECTIVES"),
    ("p", "BOARD displays (one by one):"),
    ("p", "Find the equivalent resistance of resistors in series"),
    ("p", "Find the equivalent resistance of resistors in parallel"),
    ("p", "Aadhi speaks:"),
    ("p", "In this video, we will learn how resistors combine in series and in parallel."),
    ("p", "ANIMATION: Bullets appear one by one with a soft gold glow."),
    ("h", "SCENE 1 - TITLE CARD [0:30 - 0:40]"),
    ("p", "ON SCREEN - full screen, no board, no Aadhi:"),
    ("h", "SERIES AND PARALLEL RESISTORS"),
    ("p", "ANIMATION: Title fades in over a glowing circuit board."),
    ("h", "SEGMENT 2 - HOOK: WHY DO FESTIVAL LIGHTS GO DARK? [0:40 - 1:10]"),
    ("p", "Aadhi speaks:"),
    ("p", "A string of festival lights goes completely dark when a single bulb fails."),
    ("p", "Have you ever wondered why one broken bulb switches off the whole string?"),
    ("p", "The answer lies in how the bulbs are connected."),
    ("h", "SEGMENT 3 - RESISTORS IN SERIES [1:10 - 2:30]"),
    ("p", "BOARD Title: RESISTORS IN SERIES"),
    ("p", "Aadhi speaks:"),
    ("p", "Resistors are in series when they are connected end to end, so the same current flows through each "
          "of them."),
    ("p", "The total resistance is the sum of the individual resistances."),
    ("p", "BOARD displays - formula / expression block:"),
    ("p", "R = R1 + R2 + R3"),
    ("p", "Adding a resistor in series always increases the total resistance."),
    ("p", "ANIMATION: Charges flow through three resistors in a single loop while an ammeter shows the same reading "
          "everywhere."),
    ("h", "SEGMENT 4 - WORKED EXAMPLE: SERIES CIRCUIT [2:30 - 3:20]"),
    ("p", "Aadhi speaks:"),
    ("p", "Consider three resistors of 2 ohm, 3 ohm and 5 ohm in series."),
    ("p", "Add the resistances:"),
    ("p", "R = 2 + 3 + 5"),
    ("p", "Therefore:"),
    ("p", "R = 10"),
    ("p", "The total resistance is 10 ohm."),
    ("h", "SEGMENT 5 - BRIDGE TO NEXT PART [3:20 - 3:30]"),
    ("p", "BOARD Title: COMING UP NEXT"),
    ("p", "Aadhi speaks:"),
    ("p", "Next, we will connect the same resistors side by side."),
    ("p", "NOTE: No farewell. Clean fade to black."),
    ("h", "CLIP 2 SCRIPT - PARALLEL RESISTORS"),
    ("p", "Estimated duration: 4 minutes"),
    ("h", "SEGMENT 1 - WHAT THIS VIDEO WILL COVER [0:00 - 0:20]"),
    ("p", "Aadhi speaks:"),
    ("p", "In this video, we will learn how resistors behave in parallel."),
    ("h", "SCENE 1 - TITLE CARD [0:20 - 0:30]"),
    ("p", "ON SCREEN - full screen, no board, no Aadhi:"),
    ("h", "PARALLEL RESISTORS"),
    ("h", "SEGMENT 2 - RESISTORS IN PARALLEL [0:30 - 1:40]"),
    ("p", "Aadhi speaks:"),
    ("p", "Resistors are in parallel when they are connected across the same two points, so each has the same "
          "voltage."),
    ("p", "The reciprocal of the total resistance is the sum of the reciprocals of the individual resistances."),
    ("p", "BOARD displays - formula / expression block:"),
    ("p", "1/R = 1/R1 + 1/R2 + 1/R3"),
    ("p", "Adding a resistor in parallel always decreases the total resistance."),
    ("p", "ANIMATION: Current splits into three glowing branches and joins again."),
    ("h", "SEGMENT 3 - REAL-WORLD APPLICATIONS [1:40 - 2:20]"),
    ("p", "Aadhi speaks:"),
    ("p", "Household sockets are wired in parallel so that every appliance gets the full supply voltage."),
    ("p", "Festival lights are often wired in series, which is why one failed bulb breaks the whole string."),
    ("h", "SEGMENT 4 - QUICK QUIZ [2:20 - 3:00]"),
    ("p", "Question 1:"),
    ("p", "What happens to the total resistance when a resistor is added in series?"),
    ("p", "Answer:"),
    ("p", "It increases."),
    ("p", "Question 2:"),
    ("p", "Why are household sockets wired in parallel?"),
    ("p", "Answer:"),
    ("p", "So that every appliance gets the full supply voltage."),
    ("h", "SEGMENT 5 - SUMMARY [3:00 - 3:30]"),
    ("p", "Aadhi speaks:"),
    ("p", "Series resistors add up, while parallel resistors reduce the total resistance."),
    ("p", "Thank you for watching."),
    ("p", f"Prepared by: {SME}"),
    ("p", "Word count: 640"),
]
QUESTIONS = ("What happens to the total resistance when a resistor is added in series?",
             "Why are household sockets wired in parallel?")


def sme_docx() -> bytes:
    """A two-clip SME video script with an administrative header, as a teacher would upload it."""
    import docx

    d = docx.Document()
    for line in HEADER:
        d.add_paragraph(line)
    for style, text in SCRIPT:
        if style == "h":
            d.add_heading(text, level=2)
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def assert_no_secrets(text: str, where: str) -> None:
    for secret in SECRETS:
        assert secret.lower() not in text.lower(), f"{secret!r} leaked into {where}"


# ---------------------------------------------------------------------------
# the full offline lecture through the orchestrator
# ---------------------------------------------------------------------------


def _generate(job_ctx: Any, monkeypatch: Any, data: bytes, filename: str, options: GenerationOptions) -> dict[str, Any]:
    providers = install_providers(monkeypatch)
    plans: list[Any] = []
    real = orchestrator.generate_plan

    async def spy(*args: Any, **kwargs: Any) -> Any:
        result = await real(*args, **kwargs)
        plans.append(result.plan)
        return result

    monkeypatch.setattr(orchestrator, "generate_plan", spy)
    seeded = seed_project(job_ctx.assets.storage, data, DOCX, filename, options=options.model_dump(mode="json"))
    job_ctx.project_id, job_ctx.version_id = seeded.project_id, seeded.version_id
    job_ctx.kind = "generate_lecture"
    job_ctx.payload = {"source_document_id": seeded.source_id, "options": options.model_dump(mode="json"),
                       "base_revision": 1}
    asyncio.run(orchestrator.generate_lecture(job_ctx))
    v = get_version(seeded.version_id)
    assert v.status == "ready", v.generation_meta.get("error")
    return {"version": v, "screenplay": Screenplay.model_validate(v.screenplay), "plan": plans[-1],
            "llm": providers.llm, "events": job_ctx.events}


@pytest.fixture()
def sme_lecture(job_ctx, monkeypatch, fast_audio) -> dict[str, Any]:
    return _generate(job_ctx, monkeypatch, sme_docx(), "series_parallel.docx", OPTIONS)


def test_header_values_never_reach_a_prompt_the_lecture_or_a_job_event(sme_lecture):
    llm = sme_lecture["llm"]
    schemas = {c["schema"] for c in llm.calls}
    assert {"GenBrief", "GenPlan", "GenQuiz", "GenCritique", "GenPractice"} <= schemas
    assert any(s.startswith("GenBoardScene") for s in schemas)
    for call in llm.calls:
        assert_no_secrets(call["system"] + call["prompt"], f"the {call['schema']} prompt")
        for marker in ("Estimated duration", "CLIP 1 SCRIPT", "TITLE CARD", "BRIDGE TO NEXT PART", "[0:"):
            assert marker not in call["prompt"], (marker, call["schema"])
    v, sp = sme_lecture["version"], sme_lecture["screenplay"]
    lecture = json.dumps(v.screenplay, ensure_ascii=False)
    assert_no_secrets(lecture, "the screenplay")
    for marker in PACKAGING:
        assert marker.lower() not in lecture.lower(), marker
    assert_no_secrets(json.dumps(sme_lecture["events"], ensure_ascii=False), "the job events")
    assert not [i for i in v.issues if i["code"] == "content.admin_leak"], v.issues
    excluded = v.generation_meta["ingest"]["excluded"]
    assert excluded.get("person") and excluded.get("duration") and excluded.get("timecode")
    assert v.generation_meta["ingest"]["source_format"] == "sme_script"
    # the source's own metadata fills the title-card fields; the header's people and codes do not
    assert (sp.subject_name, sp.unit_name, sp.session_number) == ("Basic Electrical Engineering", "DC Circuits", "Session 3")


# the SME's name and the video's packaging inside the narration (after the header row-wise table)
NARRATION_LEAKS = {
    "RESISTORS IN SERIES": [
        f"{SME} will now demonstrate this on the board with three resistors in a row.",
        "This video is about 4 minutes long, so keep a notebook ready while we work through it.",
        "CAMERA: zoom in on the board.",
        "NOTE TO EDITOR: insert b-roll of a festival light string here.",
        "Don't forget to like and subscribe!",
    ],
    "RESISTORS IN PARALLEL": [
        "Welcome to video 2 of the resistor networks series.",
        "As we saw in Clip 1, the same current flows through resistors in series.",
        "Hello, I am Ramesh, and I teach basic electrical engineering.",
    ],
}
NARRATION_MARKERS = ("Ramesh", "4 minutes", "CAMERA", "b-roll", "subscribe", "video 2", "Clip 1", "Hello, I am")


def sme_docx_with_narration_leaks() -> bytes:
    """The two-clip script with a row-wise header table and packaging inside the narration."""
    import docx

    d = docx.Document()
    rows = [("S.No", "Particulars", "Details")] + [(str(n), *line.split(": ", 1)) for n, line in enumerate(HEADER, 1)]
    table = d.add_table(rows=0, cols=3)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    for style, text in SCRIPT:
        (d.add_heading if style == "h" else d.add_paragraph)(text, **({"level": 2} if style == "h" else {}))
        for segment, lines in NARRATION_LEAKS.items():
            if style == "h" and segment in text:
                d.add_paragraph("Aadhi speaks:")
                for line in lines:
                    d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_names_and_packaging_inside_the_narration_never_reach_a_prompt(job_ctx, monkeypatch, fast_audio):
    out = _generate(job_ctx, monkeypatch, sme_docx_with_narration_leaks(), "series_parallel_2.docx", OPTIONS)
    calls = out["llm"].calls
    assert len(calls) >= 8
    for call in calls:
        assert_no_secrets(call["system"] + call["prompt"], f"the {call['schema']} prompt")
        for marker in NARRATION_MARKERS:
            assert marker.lower() not in call["prompt"].lower(), (marker, call["schema"])
    lecture = json.dumps(out["version"].screenplay, ensure_ascii=False)
    for marker in NARRATION_MARKERS:
        assert marker.lower() not in lecture.lower(), marker
    excluded = out["version"].generation_meta["ingest"]["excluded"]
    assert excluded.get("production_note") and excluded.get("duration") and excluded.get("person")
    # the teaching sentence after the removed packaging is still there
    plan_prompt = out["llm"].calls_for("GenPlan")[0]["prompt"]
    assert "the same current flows through resistors in series" in plan_prompt


def test_the_two_clips_become_one_continuous_lecture(sme_lecture):
    sp, plan = sme_lecture["screenplay"], sme_lecture["plan"]
    titles = [s for s in sp.scenes if s.type == "title"]
    assert len(titles) == 1 and sp.scenes[0].type == "title"  # one opening
    assert lint_intros(sp) == []  # objectives stated once, no repeated introductions
    assert not [i for i in sme_lecture["version"].issues if i["code"] == "content.duplicate_intro"]
    planned = plan.all_scenes()
    assert planned[0].narrative_role == "hook" and planned[0].bridge_in == ""
    for s in planned[1:]:
        assert s.bridge_in.strip(), f"scene {s.id} has no bridge_in"
        assert s.narrative_role, s.id
    # concepts in source order, nothing structural among them
    names = [c.title.lower() for c in plan.concept_map]
    assert any("series" in n for n in names) and any("parallel" in n for n in names)
    assert names.index(next(n for n in names if "series" in n)) < names.index(next(n for n in names if "parallel" in n))
    for bad in ("objective", "title card", "bridge", "quiz", "summary", "hook", "subject name", "what this video"):
        assert not any(bad in n for n in names), (bad, names)
    # every teaching scene opens with its bridge (the offline writer follows the plan)
    by_id = {s.id: s for s in planned}
    for scene in sp.scenes:
        if scene.type == "content":
            assert scene.beats[0].narration.startswith(by_id[scene.id].bridge_in[:30]), scene.id
    # the cold open uses the source's hook and states the objectives once
    opening = " ".join(b.narration for b in sp.scenes[0].beats)
    assert "festival lights" in opening and opening.count("you will be able to") == 1
    # the author's animation directions are visual ideas, never narration
    narration = " ".join(b.narration for s in sp.scenes for b in s.all_beats())
    assert "glow" not in narration.lower() and "ammeter shows the same reading" not in narration
    assert any(s.visual_note_ids for s in planned)


def test_quizzes_use_the_sources_own_questions(sme_lecture):
    quizzes = [s for s in sme_lecture["screenplay"].scenes if isinstance(s, QuizScene)]
    assert quizzes
    asked = {q.question for q in quizzes}
    assert asked & set(QUESTIONS), asked
    quiz = next(q for q in quizzes if q.question == QUESTIONS[0])
    assert quiz.options[quiz.correct_index] == "It increases"


def test_the_real_sme_template_becomes_a_coherent_lecture(job_ctx, monkeypatch, fast_audio):
    out = _generate(job_ctx, monkeypatch, sample_template_bytes(), "session2.docx",
                    GenerationOptions(target_minutes=15, quiz_every_n_concepts=2))
    sp, plan, v = out["screenplay"], out["plan"], out["version"]
    assert [s.type for s in sp.scenes].count("title") == 1 and lint_intros(sp) == []
    assert not [i for i in v.issues if i["code"].startswith("content.")], v.issues
    assert not [i for i in v.issues if i["source"] == "lint" and i["severity"] == "error"], v.issues
    assert all(s.bridge_in for s in plan.all_scenes()[1:])
    names = [c.title for c in plan.concept_map]
    assert "Boolean postulates" in names and "De Morgan's theorems" in names and "Boolean minimization" in names
    assert not any(re.search(r"(?i)objective|quiz|summary|hook|secret language", n) for n in names), names
    questions = {s.question for s in sp.scenes if isinstance(s, QuizScene)}
    assert "Who developed Boolean Algebra?" in questions and "What is A + 0 equal to?" in questions
    narration = " ".join(b.narration for s in sp.scenes for b in s.all_beats())
    for marker in ("In this video", "Thank you for watching", "Let's summarize", "glow"):
        assert marker.lower() not in narration.lower(), marker
    simulations = [s for s in sp.scenes if s.type == "simulation"]
    if simulations and simulations[0].manim.template == "equation_steps":  # the source's own derivation
        steps = [x["latex"] for x in simulations[0].manim.params["steps"]]
        assert steps[0].startswith("Y = AB") and steps[-1] == "Y = A"


# ---------------------------------------------------------------------------
# fake_content: reading a source like a teacher (no I/O)
# ---------------------------------------------------------------------------


def test_prose_sentences_join_script_fragments_and_drop_recording_talk():
    text = "\n".join([
        "Unlike ordinary algebra, Boolean Algebra deals with only two values:", "0 and 1", "Where:", "0 means FALSE",
        "1 means TRUE", "A student can enter an examination hall only if:", "ID Card is available", "AND",
        "Hall Ticket is available.", "AND Operation", "Represented by a dot.", "In this video, we will learn:",
        "Thank you for watching.", "**A = 0 or 1**", "Within the elastic limit, stress is directly proportional to strain:",
    ])
    assert fake_content.prose_sentences(text) == [
        "Unlike ordinary algebra, Boolean Algebra deals with only two values: 0 and 1.",
        "Where: 0 means FALSE, 1 means TRUE.",
        "A student can enter an examination hall only if: ID Card is available AND Hall Ticket is available.",
        "AND Operation is represented by a dot.",
        "Within the elastic limit, stress is directly proportional to strain.",
    ]
    assert fake_content.prose_sentences("A lecture about video coding.\nIn this video codec, frames are predicted.")


def test_section_roles_and_concept_names():
    kinds = {h: fake_content.section_kind(h) for h in (
        "Clip > Learning objectives", "SEGMENT 1 - WHAT THIS VIDEO WILL COVER", "SCENE 1 - TITLE CARD", "Bridge to next part",
        "QUICK QUIZ", "Practice", "SUMMARY", "HOOK: THE SECRET LANGUAGE", "EXAMPLE 1: BOOLEAN MINIMIZATION",
        "REAL-WORLD ANALOGY", "Real-world engineering applications", "6. A common misconception", "V = I × R",
        "Introduction to Boolean laws", "Commutative law", "Test of hypotheses",
    )}
    assert kinds == {
        "Clip > Learning objectives": "structural", "SEGMENT 1 - WHAT THIS VIDEO WILL COVER": "structural",
        "SCENE 1 - TITLE CARD": "structural", "Bridge to next part": "structural", "QUICK QUIZ": "quiz",
        "Practice": "quiz", "SUMMARY": "structural", "HOOK: THE SECRET LANGUAGE": "hook",
        "EXAMPLE 1: BOOLEAN MINIMIZATION": "example", "REAL-WORLD ANALOGY": "analogy",
        "Real-world engineering applications": "application", "6. A common misconception": "aside", "V = I × R": "aside",
        "Introduction to Boolean laws": "intro", "Commutative law": "content", "Test of hypotheses": "content",
    }
    assert fake_content.concept_name("POSTULATE 1: CLOSURE PROPERTY") == "Closure Property"
    assert fake_content.concept_name("WHAT IS BOOLEAN ALGEBRA?") == "Boolean Algebra"
    assert fake_content.concept_name("BOOLEAN POSTULATES INTRODUCTION") == "Boolean Postulates"
    assert fake_content.natural_case("Fundamental Boolean Laws", "these laws ... Boolean algebra") == "Fundamental Boolean laws"
    assert fake_content.focus_case("Identity elements", "the identity of A") == "identity elements"


def test_questions_formulas_and_tables():
    quiz = "Question 1:\nWhat values can a Boolean variable take?\nAnswer:\n0 and 1.\n- Q2: Who developed it? Answer: George Boole."
    assert fake_content.qa_pairs(quiz) == [("What values can a Boolean variable take?", "0 and 1."),
                                           ("Who developed it?", "George Boole.")]
    assert fake_content.qa_pairs("But have you ever wondered why?", any_question=False) == []
    assert fake_content.split_question("What is A + 0 equal to? Answer: A.") == ("What is A + 0 equal to?", "A")
    assert fake_content.find_formula("Let: A = ID Card") is None  # a letter naming a thing
    assert fake_content.find_formula("Y = AB + AB̅") is None  # the expression goes on
    assert fake_content.find_formula("I = V / R = 12 / 4 = 3 A") == ("I", "V", "/", "R")
    assert fake_content.latex_of("(A · B)̅ = A̅ + B̅") == "\\overline{(A \\cdot B)} = \\overline{A} + \\overline{B}"
    assert fake_content.latex_of("A = 0 or 1") is None and fake_content.latex_of("A = ID Card") is None
    chain = fake_content.derivation("Y = AB + AB̅\nFactor A:\nY = A(B + B̅)\nSince:\nB + B̅ = 1\nY = A × 1\nY = A")
    assert [e for e, _ in chain] == ["Y = AB + AB̅", "Y = A(B + B̅)", "Y = A × 1", "Y = A"]
    assert chain[1][1] == "Factor A" and "B + B̅ = 1" in chain[2][1]
    table = "| Gate | Boolean expression |\n|---|---|\n| AND | Y = A · B |\n| OR | Y = A + B |"
    assert fake_content.table_sentences(table) == ["AND: Boolean expression Y = A · B.", "OR: Boolean expression Y = A + B."]


def _template_chunks() -> list[dict[str, Any]]:
    scope = scope_source(docx_to_markdown(sample_template_bytes()).markdown)
    return [{"id": c.id, "heading": " > ".join(c.heading_path), "text": c.text} for c in chunk_markdown(scope.markdown)]


def test_concepts_of_the_sme_template_are_its_teaching_ideas():
    found = fake_content.source_concepts(_template_chunks())
    names = [c.name for c in found]
    assert names == ["Boolean Algebra", "Boolean Variables", "Three Basic Logical Operations", "Boolean Postulates",
                     "Fundamental Boolean Laws", "De Morgan's Theorems", "Boolean Minimization"]
    laws = found[4]
    assert [p for p, _ in laws.parts] == ["Commutative Law", "Associative Law", "Distributive Law", "Idempotent Law",
                                          "Null Law"]
    assert laws.context  # "Introduction to Boolean laws" leads into the laws without being a concept of its own
    assert found[6].examples and found[6].applications and found[2].analogies


def test_offline_brief_for_the_sme_template_is_valid_and_teaching_only():
    from aadhi.pipeline.base import IngestResult, SourceChunk

    chunks = _template_chunks()
    ingest = IngestResult(markdown="x", source_format="sme_script",
                          chunks=[SourceChunk(id=c["id"], heading_path=c["heading"].split(" > "), text=c["text"])
                                  for c in chunks])
    _, user = build_brief_prompt(ingest, OPTIONS)
    gen = GenBrief.model_validate(fake_content.brief_responder(user, GenBrief))
    assert brief_problems(gen, {c["id"] for c in chunks}) == ([], [])  # no re-ask needed
    assert "Who developed Boolean Algebra? Answer: George Boole." in [q.text for q in gen.source_questions]
    laws = next(c for c in gen.concepts if c.key == "fundamental_boolean_laws")
    assert any("A + B = B + A" in f.text for f in laws.key_facts)
    assert laws.must_explain[0].startswith("Commutative law: Order does not matter")
    text = gen.model_dump_json()
    for marker in ("In this video", "Thank you", "Learning objectives", "Quick quiz"):
        assert marker.lower() not in text.lower(), marker


def test_offline_plan_follows_the_brief_with_roles_bridges_and_source_questions(templates):
    from aadhi.pipeline.prompting import json_section

    chunks = _template_chunks()
    brief_user = json_section("Request", {"depth": "standard"}) + json_section("Source chunks", chunks)
    brief = fake_content.brief_responder(brief_user, GenBrief)
    request = {"target_seconds": 900, "quiz_every_n_concepts": 2, "include_quizzes": True,
               "allowed_scene_types": ["title", "content", "example", "summary", "chapter_card", "recap",
                                       "quiz_checkpoint", "simulation"],
               "session_title": "Boolean postulates, laws and minimization"}
    prompt = (json_section("Request", request) + json_section("Concepts to teach", brief)
              + json_section("Animation templates", [{"name": "equation_steps"}]) + json_section("Source chunks", chunks))
    gen = GenPlan.model_validate(fake_content.plan_responder(prompt, GenPlan))
    assert [c.key for c in gen.concept_map] == brief["teaching_order"]
    scenes = [s for ch in gen.chapters for s in ch.scenes]
    assert scenes[0].type == "title" and scenes[0].narrative_role == "hook" and not scenes[0].bridge_in
    assert [s.type for s in scenes].count("title") == 1
    assert all(s.bridge_in and s.narrative_role for s in scenes[1:])
    assert scenes[0].source_refs == ["c0002"]  # the hook section opens the lecture
    quizzes = [s for s in scenes if s.type == "quiz_checkpoint"]
    assert any(k.startswith("Who developed Boolean Algebra?") for q in quizzes for k in q.key_points)
    sims = [s for s in scenes if s.type == "simulation"]
    assert len(sims) == 1 and sims[0].manim_template == "equation_steps"
    assert extract_json(prompt, "Concepts to teach") == brief
