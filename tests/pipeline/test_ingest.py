"""ingest: PDF (subprocess worker), DOCX, TXT/MD, caching, limits, detection heuristics."""

from __future__ import annotations

import asyncio

import pytest

from aadhi.pipeline import ingest as ingest_mod
from aadhi.pipeline.chunking import chunk_markdown, detect_language, truncate_markdown
from aadhi.pipeline.docx_extract import DocxRejected, check_zip, docx_to_markdown
from aadhi.pipeline.ingest import IngestError, decode_text, ingest_source, source_kind, text_to_markdown
from tests.pipeline.dbutil import asset_refs, get_source, seed_project
from tests.pipeline.sources import make_docx, make_pdf

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def run_ingest(job_ctx, data: bytes, mime: str, filename: str):
    seeded = seed_project(job_ctx.assets.storage, data, mime, filename)
    job_ctx.project_id = seeded.project_id
    source = get_source(seeded.source_id)
    return seeded, asyncio.run(ingest_source(job_ctx, source))


def test_pdf_ingest_end_to_end(job_ctx):
    seeded, res = run_ingest(job_ctx, make_pdf(), PDF, "notes.pdf")
    md = res.markdown
    assert res.pages == 3
    for n in (1, 2, 3):
        assert f"<!-- page {n} -->" in md
    assert "# 1. Voltage and current" in md and "# 2. Resistance" in md
    assert "Basic Electrical Engineering - Unit 2" not in md  # running header removed
    assert "Page 2" not in md  # page numbers removed
    assert "| Quantity | Symbol | Unit |" in md and "| Resistance | R | ohm |" in md
    assert len(res.figures) == 1
    fig = res.figures[0]
    assert fig.id == "fig-p2-1" and fig.page == 2 and fig.caption.startswith("Figure 1")
    assert f"[Figure {fig.id}: Figure 1" in md
    assert fig.asset_key and fig.asset_key.startswith("figure-") and (fig.width, fig.height) == (320, 200)
    asset = job_ctx.assets.get(fig.asset_key)
    assert asset.kind == "figure" and asset.mime == "image/png"
    assert [c.id for c in res.chunks][:2] == ["c0001", "c0002"]
    assert {c.page for c in res.chunks} == {1, 2, 3}
    assert res.chunks[0].heading_path == ["1. Voltage and current"]
    assert res.detected_language == "en-IN"
    assert not res.attach_original
    src = get_source(seeded.source_id)
    assert src.extracted_key and src.extracted_key.startswith("private/extract/") and src.page_count == 3
    refs = asset_refs(seeded.project_id)
    assert fig.asset_key in refs and any(r.startswith("extract-") for r in refs)


def test_pdf_ingest_is_cached(job_ctx, monkeypatch):
    data = make_pdf()
    _, first = run_ingest(job_ctx, data, PDF, "a.pdf")

    async def boom(*a, **k):
        raise AssertionError("worker must not run on a cache hit")

    monkeypatch.setattr(ingest_mod, "run_pdf_worker", boom)
    seeded = seed_project(job_ctx.assets.storage, data + b"\n%", PDF, "b.pdf")  # different project, other blob
    src = get_source(seeded.source_id)
    src.sha256 = get_source(1).sha256  # same content hash as the first upload
    second = asyncio.run(ingest_source(job_ctx, src))
    assert second.markdown == first.markdown
    assert second.source_storage_key == src.storage_key
    assert first.figures[0].asset_key in asset_refs(seeded.project_id)


def test_pdf_page_cap_and_scanned_detection(job_ctx, monkeypatch):
    import pymupdf

    from tests.pipeline.fakes import png_bytes

    doc = pymupdf.open()
    for _ in range(4):  # image-only pages look scanned
        page = doc.new_page()
        page.insert_image(page.rect, stream=png_bytes(600, 800, "scan"))
    data = doc.tobytes()
    monkeypatch.setattr(job_ctx.settings, "max_source_pages", 2)
    _, res = run_ingest(job_ctx, data, PDF, "scan.pdf")
    assert res.pages == 2 and res.truncated and res.attach_original
    assert any("scanned" in w for w in res.warnings) and any("first 2 of 4" in w for w in res.warnings)


def test_maths_heavy_pdf_attaches_original(job_ctx):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 80
    for _ in range(6):  # equations typeset in a maths (Symbol) font between short prose lines
        page.insert_text((72, y), "We now use the following relation:", fontsize=11, fontname="helv")
        page.insert_text((72, y + 18), "a + b = g d (p r) / q + s t", fontsize=11, fontname="symb")
        y += 40
    _, res = run_ingest(job_ctx, doc.tobytes(), PDF, "maths.pdf")
    assert res.attach_original and any("maths-heavy" in w for w in res.warnings)


def test_corrupt_and_encrypted_pdf(job_ctx):
    import pymupdf

    with pytest.raises(IngestError, match="cannot read"):
        run_ingest(job_ctx, b"%PDF-1.4 not really a pdf", PDF, "bad.pdf")
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")
    with pytest.raises(IngestError, match="password"):
        run_ingest(job_ctx, data, PDF, "locked.pdf")


def test_pdf_worker_timeout_kills_process(job_ctx, monkeypatch):
    monkeypatch.setattr(ingest_mod, "pdf_timeout_seconds", lambda pages: 0.01)
    with pytest.raises(IngestError, match="longer than"):
        run_ingest(job_ctx, make_pdf(), PDF, "slow.pdf")


def test_docx_ingest(job_ctx):
    seeded, res = run_ingest(job_ctx, make_docx(), DOCX, "notes.docx")
    md = res.markdown
    assert "# Ohm's Law" in md and "## Resistance" in md
    assert "Resistance is the **opposition** a material offers to current." in md
    assert "- Copper has low resistance" in md
    assert "| Quantity | Unit |" in md and "| Resistance | ohm |" in md
    assert len(res.figures) == 1 and res.figures[0].id == "fig-1"
    assert res.figures[0].caption == "Figure 2: Resistors in a circuit"
    assert "[Figure fig-1: Figure 2: Resistors in a circuit]" in md
    assert res.figures[0].asset_key in asset_refs(seeded.project_id)
    assert res.chunks and res.chunks[0].heading_path[0] == "Ohm's Law"


def test_docx_zip_checks():
    with pytest.raises(DocxRejected):
        check_zip(b"PK not a zip")
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "x")
    with pytest.raises(DocxRejected, match="Word"):
        docx_to_markdown(buf.getvalue())


def test_text_and_markdown_ingest(job_ctx):
    txt = "UNIT 2 Electric circuits\n\n1. Introduction\nCurrent flows when a circuit is closed.\n\n1.1 Ohm's law\nV equals I R.\n"
    _, res = run_ingest(job_ctx, txt.encode("utf-8"), "text/plain", "notes.txt")
    assert "# UNIT 2 Electric circuits" in res.markdown and "## 1.1 Ohm's law" in res.markdown
    _, md = run_ingest(job_ctx, b"# Title\n\nSome text about resistors and current.\n", "text/markdown", "n.md")
    assert md.chunks[0].heading_path == ["Title"] and md.pages is None


def test_tamil_detection_and_decoding():
    assert detect_language("மின்னோட்டம் என்பது மின்னூட்டத்தின் ஓட்டம். Current") == "ta-IN"
    assert detect_language("वोल्टेज और धारा") == "hi-IN"
    assert detect_language("plain english") == "en-IN"
    assert detect_language("1234 !!") is None
    assert decode_text("﻿hello".encode()) == "hello"
    assert decode_text("héllo".encode("cp1252")) == "héllo"
    assert decode_text("hi".encode("utf-16")) == "hi"
    assert source_kind("application/octet-stream", "x.DOCX") == "docx"


def test_truncation_and_chunking_limits(job_ctx, monkeypatch):
    para = "Resistance opposes current in every conductor we use. " * 40
    md = "\n\n".join(f"## Section {i}\n\n{para}" for i in range(20))
    cut, truncated = truncate_markdown(md, 5000)
    assert truncated and len(cut) <= 5001
    chunks = chunk_markdown(md)
    assert all(len(c.text) <= 2000 for c in chunks)
    assert len({c.id for c in chunks}) == len(chunks)
    monkeypatch.setattr(job_ctx.settings, "max_source_chars", 4000)
    _, res = run_ingest(job_ctx, md.encode(), "text/markdown", "long.md")
    assert res.truncated and any("very long" in w for w in res.warnings)


def test_empty_source_rejected(job_ctx):
    with pytest.raises(IngestError, match="no readable text"):
        run_ingest(job_ctx, b"   \n ", "text/plain", "empty.txt")


def test_text_to_markdown_headings():
    md = text_to_markdown("Chapter 3: Motors\r\nBody text here.\r\n2.1 DC motors\r\nMore text.")
    assert md.splitlines()[0] == "# Chapter 3: Motors" and "## 2.1 DC motors" in md


def test_pseudo_headings_in_unstyled_docx_and_text():
    import io

    import docx

    from aadhi.pipeline.chunking import chunk_markdown, looks_like_heading
    from aadhi.pipeline.docx_extract import docx_to_markdown

    assert looks_like_heading("SEGMENT 1 - WHAT THIS VIDEO WILL COVER")
    assert not looks_like_heading("DETAILED ANIMATION / VISUAL:")  # a label, not a title
    assert not looks_like_heading("Resistance opposes current in every conductor we use.")
    assert looks_like_heading("Kirchhoff's voltage law", bold=True)
    assert not looks_like_heading("OK")

    d = docx.Document()  # no heading styles at all, like many teacher-written scripts
    d.add_paragraph("BOOLEAN ALGEBRA BASICS")
    d.add_paragraph("Boolean algebra works with two values, true and false.")
    p = d.add_paragraph()
    p.add_run("De Morgan's theorem").bold = True
    d.add_paragraph("The complement of a product equals the sum of the complements.")
    buf = io.BytesIO()
    d.save(buf)
    md = docx_to_markdown(buf.getvalue()).markdown
    assert "## BOOLEAN ALGEBRA BASICS" in md and "### De Morgan's theorem" in md
    chunks = chunk_markdown(md)
    assert chunks[-1].heading_path == ["BOOLEAN ALGEBRA BASICS", "De Morgan's theorem"]

    styled = docx.Document()  # with real heading styles, look-alike paragraphs become minor sub-headings
    styled.add_heading("Logic gates", level=1)
    styled.add_paragraph("AND GATE EXAMPLES")
    styled.add_paragraph("An AND gate outputs one only when every input is one.")
    buf = io.BytesIO()
    styled.save(buf)
    assert "#### AND GATE EXAMPLES" in docx_to_markdown(buf.getvalue()).markdown

    md = text_to_markdown("INTRODUCTION TO DIODES\nA diode conducts in one direction.\nNOTE:\nkeep it cool.")
    assert md.splitlines()[0] == "## INTRODUCTION TO DIODES" and "## NOTE:" not in md


# ---------------------------------------------------------------------------
# regressions (review round 2)
# ---------------------------------------------------------------------------


def test_readable_chars_ignores_markers_and_markup():
    from aadhi.pipeline.ingest import readable_chars

    assert readable_chars("<!-- page 1 -->\n[Figure fig-p1-1]\n\n<!-- page 2 -->\n[Figure fig-p2-1: Figure 2: a]\n") == 0
    assert readable_chars("# Ohm\n\n| a |\n|---|\n") == len("Ohma")
    assert readable_chars("V = IR") == len("V=IR")


def test_image_only_sources_are_not_readable_text(job_ctx):
    import io

    import docx

    from tests.pipeline.fakes import png_bytes

    d = docx.Document()  # a Word file holding only a picture (e.g. a pasted scan)
    d.add_picture(io.BytesIO(png_bytes(200, 150, "scan")))
    buf = io.BytesIO()
    d.save(buf)
    with pytest.raises(IngestError, match="no readable text"):
        run_ingest(job_ctx, buf.getvalue(), DOCX, "scan.docx")
    with pytest.raises(IngestError, match="no readable text"):
        run_ingest(job_ctx, b"<!-- page 1 -->\n[Figure fig-1]\n", "text/markdown", "only-figure.md")


def test_scanned_pdf_warnings_do_not_promise_an_attachment(job_ctx):
    import pymupdf

    from tests.pipeline.fakes import png_bytes

    doc = pymupdf.open()
    for _ in range(3):
        page = doc.new_page()
        page.insert_image(page.rect, stream=png_bytes(600, 800, "scan-page"))
    _, res = run_ingest(job_ctx, doc.tobytes(), PDF, "scan.pdf")
    assert res.attach_original
    assert any("could not be extracted" in w for w in res.warnings)
    assert not any("given to the model" in w for w in res.warnings)  # decided (and reported) at planning time
